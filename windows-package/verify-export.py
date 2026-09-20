import uuid
import os,tempfile,subprocess,urllib.request,urllib.error,json,socket,time,pathlib,shutil,datetime
from openpyxl import load_workbook
package=pathlib.Path(__file__).resolve().parent/'release'/'RDIMM-Windows'
node='/Users/yanmeng/.cache/codex-runtimes/codex-primary-runtime/dependencies/node/bin/node'
opener=urllib.request.build_opener(urllib.request.ProxyHandler({}))
with tempfile.TemporaryDirectory(prefix='rdimm-export-') as tmp:
 root=pathlib.Path(tmp);shutil.copy(package/'server.cjs',root);shutil.copy(package/'initial-data.json',root)
 with socket.socket() as s:s.bind(('127.0.0.1',0));port=s.getsockname()[1]
 base=f'http://127.0.0.1:{port}'
 p=subprocess.Popen([node,'--no-warnings','server.cjs'],cwd=root,env={**os.environ,'RDIMM_PORT':str(port)},stdout=subprocess.PIPE,stderr=subprocess.STDOUT)
 def post(endpoint,data,origin=base):
  req=urllib.request.Request(base+endpoint,data=json.dumps(data).encode(),headers={'Content-Type':'application/json','Origin':origin})
  try:
   with opener.open(req) as r:return r.status,json.load(r)
  except urllib.error.HTTPError as e:return e.code,json.load(e)
 try:
  for _ in range(60):
   try:opener.open(base+'/health');break
   except:time.sleep(.1)
  assert post('/api/export',{},'http://evil.example')[0]==403
  assert post('/api/shutdown',{})[0]==400
  code,created=post('/api/samples',dict(action='create',customer='=1+2',owner='中文申请人',spec='32G',batch='000123',quantity=12,note='换行\n<&>'));assert code==200
  snapshot=json.load(opener.open(base+'/api/samples'))
  code,result=post('/api/export',{});assert code==200,result
  f=root/result['path'];wb=load_workbook(f);assert wb.sheetnames==['送样记录','归还明细']
  ws=wb.worksheets[0];assert ws.max_row==len(snapshot['rows'])+4
  records={r[0].value:r for r in list(ws.rows)[4:]};row=records[created['id']]
  assert row[1].value=='=1+2' and row[1].data_type=='s'
  assert row[4].value=='000123' and row[4].data_type=='s'
  assert row[5].value==12 and row[6].value==0 and row[8].value==0
  assert row[15].value=='换行\n<&>'
  assert ws.freeze_panes=='A5';assert ws.auto_filter.ref.endswith(str(ws.max_row))
  assert isinstance(row[9].value,datetime.datetime)
  assert wb.worksheets[1].max_row==len(snapshot['returns'])+4
  before=f.read_bytes();code,second=post('/api/export',{});assert code==200 and second['filename']!=result['filename'];assert f.read_bytes()==before
  shutil.copyfile(f,package.parent.parent.parent/'.work/export-verified.xlsx')
  saved=root/'saved';(root/'back_up').rename(saved);(root/'back_up').write_text('blocked directory')
  before_reset=json.load(opener.open(base+'/api/samples'))
  reset_key=str(uuid.uuid4())
  assert post('/api/samples',{'action':'reset','confirm':True,'requestId':reset_key})[0]==500
  assert json.load(opener.open(base+'/api/samples'))==before_reset
  assert post('/api/shutdown',{'confirm':True,'export':True})[0]==500
  assert opener.open(base+'/health').status==200
  (root/'back_up').unlink();saved.rename(root/'back_up')
  code,reset=post('/api/samples',{'action':'reset','confirm':True,'requestId':reset_key});assert code==200,reset
  backup=load_workbook(root/reset['excelBackup'])
  assert backup.worksheets[0].max_row==len(before_reset['rows'])+4
  assert backup.worksheets[1].max_row==len(before_reset['returns'])+4
  assert json.load(opener.open(base+'/api/samples'))=={'rows':[],'returns':[]}
  count=len(list((root/'back_up').glob('*.xlsx')))
  assert post('/api/samples',{'action':'reset','confirm':True,'requestId':reset_key})==(200,reset)
  assert len(list((root/'back_up').glob('*.xlsx')))==count
  code,empty=post('/api/export',{});assert code==200
  empty_wb=load_workbook(root/empty['path']);assert all(s.max_row==4 for s in empty_wb.worksheets)
  code,result=post('/api/shutdown',{'confirm':True,'export':True});assert code==200;assert (root/result['path']).is_file()
  p.wait(timeout=5);assert p.returncode==0
  print('PASS: all records, returns, numeric/date/text cells, collision protection, origin checks, failed-export keeps app running, export then exit')
 finally:
  if p.poll() is None:p.terminate();p.wait(timeout=5)
