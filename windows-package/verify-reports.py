from openpyxl import load_workbook
import os,tempfile,subprocess,urllib.request,urllib.error,json,socket,time,pathlib,shutil,datetime,uuid,sqlite3
package=pathlib.Path(__file__).resolve().parent/'release'/'RDIMM-Windows'
node='/Users/yanmeng/.cache/codex-runtimes/codex-primary-runtime/dependencies/node/bin/node'
opener=urllib.request.build_opener(urllib.request.ProxyHandler({}))
with tempfile.TemporaryDirectory(prefix='rdimm-reports-') as tmp:
 root=pathlib.Path(tmp);shutil.copy(package/'server.cjs',root);shutil.copy(package/'initial-data.json',root)
 with socket.socket() as s:s.bind(('127.0.0.1',0));port=s.getsockname()[1]
 base=f'http://127.0.0.1:{port}';env={**os.environ,'RDIMM_PORT':str(port)}
 def launch():
  p=subprocess.Popen([node,'--no-warnings','server.cjs'],cwd=root,env=env,stdout=subprocess.PIPE,stderr=subprocess.STDOUT)
  for _ in range(60):
   try:opener.open(base+'/health');return p
   except:time.sleep(.1)
  raise RuntimeError(p.stdout.read().decode())
 def send(url,data,ctype='application/json',origin=base):
  try:
   with opener.open(urllib.request.Request(base+url,data=data,headers={'Content-Type':ctype,'Origin':origin})) as r:return r.status,json.load(r)
  except urllib.error.HTTPError as e:return e.code,json.load(e)
 def post(**data):return send('/api/samples',json.dumps(data).encode())
 def upload(values,filename=None,content=b'pdf',origin=base):
  boundary='rdimm'+uuid.uuid4().hex;parts=[]
  for k,v in values.items():parts.append(f'--{boundary}\r\nContent-Disposition: form-data; name="{k}"\r\n\r\n{v}\r\n'.encode())
  if filename:parts.append(f'--{boundary}\r\nContent-Disposition: form-data; name="report_file"; filename="{filename}"\r\nContent-Type: application/octet-stream\r\n\r\n'.encode()+content+b'\r\n')
  parts.append(f'--{boundary}--\r\n'.encode());return send('/api/returns',b''.join(parts),'multipart/form-data; boundary='+boundary,origin)
 def read():return json.load(opener.open(base+'/api/samples'))
 p=launch()
 try:
  code,row=post(action='create',customer='报告测试',owner='测试员',spec='64G',quantity=10);assert code==200;sid=row['id']
  day=str(datetime.datetime.now(datetime.timezone(datetime.timedelta(hours=8))).date())
  assert post(action='send',id=sid,batch='00123',sent=day,due=day)[0]==200
  key=str(uuid.uuid4());values=dict(action='return',id=sid,requestId=key,quantity=3,returned_on=day,report_provided='yes')
  before=read();assert upload(values,'bad.exe')[0]==400;assert read()==before
  assert upload(values,'report.pdf',origin='http://evil.example')[0]==403
  assert upload({**values,'report_provided':'no'},'report.pdf')[0]==400
  pdf=b'%PDF-1.4\nreport bytes';code,result=upload(values,'客户报告.pdf',pdf);assert code==200,result
  assert any(f.read_bytes()==pdf for f in (root/'Attachment').rglob('*') if f.is_file())
  record=next(r for r in read()['returns'] if r['id']==key);assert record['report_provided']=='yes' and record['report_name']=='客户报告.pdf'
  response=opener.open(base+'/api/reports/'+key);assert response.read()==pdf;assert response.headers['Content-Type']=='application/pdf';assert response.headers['Content-Disposition'].startswith('inline;')
  assert upload(values,'客户报告.pdf',pdf)[0]==409
  assert next(r for r in read()['rows'] if r['id']==sid)['returned']==3
  second=str(uuid.uuid4());assert upload({**values,'requestId':second,'report_provided':'no'})[0]==200
  assert upload(dict(action='report',requestId=second,report_provided='yes'),'测试.xlsx',b'xlsx')[0]==200
  response=opener.open(base+'/api/reports/'+second);assert response.read()==b'xlsx';assert response.headers['Content-Disposition'].startswith('attachment;')
  p.terminate();p.wait();shutil.rmtree(root/'Attachment');p=launch();assert opener.open(base+'/api/reports/'+key).read()==pdf
  assert any(f.read_bytes()==pdf for f in (root/'Attachment').rglob('*') if f.is_file())
  (root/'Attachment').rename(root/'saved-attachments');(root/'Attachment').write_text('blocked')
  before=read();assert upload({**values,'requestId':str(uuid.uuid4())},'report.pdf',pdf)[0]==500;assert read()==before
  (root/'Attachment').unlink();(root/'saved-attachments').rename(root/'Attachment')
  # Failure inserting file must roll back the associated return and quantity.
  db=root/'data/rdimm.sqlite'
  with sqlite3.connect(db) as c:c.execute("CREATE TRIGGER reject_report BEFORE INSERT ON report_files BEGIN SELECT RAISE(ABORT,'test'); END")
  before=read();assert upload({**values,'requestId':str(uuid.uuid4())},'report.pdf',pdf)[0]==500;assert read()==before
  with sqlite3.connect(db) as c:c.execute('DROP TRIGGER reject_report')
  code,result=post(action='reset',confirm=True,requestId=str(uuid.uuid4()));assert code==200,result
  book=load_workbook(root/result['excelBackup']);sheet=book['归还明细'];exported={r[0]:r for r in sheet.iter_rows(min_row=5,values_only=True)}
  assert exported[key][-2:]==('已提供','客户报告.pdf')
  with sqlite3.connect(root/result['backup']) as c:assert c.execute('SELECT data FROM report_files WHERE return_id=?',(key,)).fetchone()[0]==pdf
  with sqlite3.connect(db) as c:assert c.execute('SELECT COUNT(*) FROM report_files').fetchone()[0]==0
  assert any(f.read_bytes()==pdf for f in (root/'Attachment').rglob('*') if f.is_file())
  try:opener.open(base+'/api/reports/'+key);raise AssertionError('deleted report still visible')
  except urllib.error.HTTPError as e:assert e.code==404
  print('PASS: multipart report upload; invalid files and origin rejected; atomic return/file save; duplicate protection; supplement; inline/download; restart; backup includes attachment; reset cascade')
 finally:
  if p.poll() is None:p.terminate();p.wait()
